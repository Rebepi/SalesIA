import io
import json
from decimal import Decimal
from uuid import UUID, uuid4
from datetime import datetime, timezone, timedelta
import pytest
from httpx import AsyncClient, ASGITransport
from sqlalchemy import select
from openpyxl import load_workbook

from salesia.main import app
from salesia.shared.infrastructure.database import sync_engine, sync_session_factory, Base
from salesia.shared.infrastructure.models import (
    CompanyModel,
    BranchModel,
    CustomerModel,
    ProductModel,
    SaleModel,
    SaleLineModel,
    KardexMovementModel,
    WarehouseModel,
    UserModel
)
from salesia.shared.infrastructure.transactions import RefundModel
from salesia.modules.operaciones.infrastructure_models import AnalyticsArtifact, SalesRepresentative
from tests.support.auth_helpers import create_durable_identity

COMPANY_A_ID = UUID("00000000-0000-0000-0000-000000000001")
COMPANY_B_ID = UUID("00000000-0000-0000-0000-000000000009")
BRANCH_A_ID = UUID("00000000-0000-0000-0000-000000000002")
USER_A_ID = UUID("00000000-0000-0000-0000-000000000003")
SUSPENDED_USER_ID = UUID("00000000-0000-0000-0000-000000000033")
CUSTOMER_A_ID = UUID("00000000-0000-0000-0000-000000000004")
PRODUCT_A_ID = UUID("00000000-0000-0000-0000-000000000005")
WAREHOUSE_A_ID = UUID("00000000-0000-0000-0000-000000000006")

@pytest.fixture(autouse=True, scope="module")
def setup_reporting_complete_db():
    Base.metadata.create_all(sync_engine)
    with sync_session_factory() as session:
        comp_a = session.execute(
            select(CompanyModel).where(CompanyModel.id == COMPANY_A_ID)
        ).scalar_one_or_none()
        if not comp_a:
            comp_a = CompanyModel(
                id=COMPANY_A_ID,
                tax_id="20999999999",
                legal_name="EMPRESA COMERCIAL MATRIZ S.A.C.",
                trade_name="SALESIA CORP",
                base_currency="USD",
                is_active=True
            )
            session.add(comp_a)

        comp_b = session.execute(
            select(CompanyModel).where(CompanyModel.id == COMPANY_B_ID)
        ).scalar_one_or_none()
        if not comp_b:
            comp_b = CompanyModel(
                id=COMPANY_B_ID,
                tax_id="20888888888",
                legal_name="EMPRESA COMPETIDORA B S.A.C.",
                trade_name="CORP B",
                base_currency="USD",
                is_active=True
            )
            session.add(comp_b)

        session.flush()
        branch = session.execute(
            select(BranchModel).where(BranchModel.id == BRANCH_A_ID)
        ).scalar_one_or_none()
        if not branch:
            branch = BranchModel(
                id=BRANCH_A_ID,
                company_id=COMPANY_A_ID,
                code="SUC-01",
                name="Sucursal Principal Lima",
                timezone="America/Lima",
                is_active=True
            )
            session.add(branch)

        cust = session.execute(
            select(CustomerModel).where(CustomerModel.id == CUSTOMER_A_ID)
        ).scalar_one_or_none()
        if not cust:
            cust = CustomerModel(
                id=CUSTOMER_A_ID,
                company_id=COMPANY_A_ID,
                code="CLI-9999",
                doc_type="RUC",
                doc_number="00020555555",
                legal_name="CLIENTE CORPORATIVO S.A.C.",
                credit_limit=Decimal("50000.00"),
                current_balance=Decimal("0.00"),
                is_active=True
            )
            session.add(cust)

        prod = session.execute(
            select(ProductModel).where(ProductModel.id == PRODUCT_A_ID)
        ).scalar_one_or_none()
        if not prod:
            prod = ProductModel(
                id=PRODUCT_A_ID,
                company_id=COMPANY_A_ID,
                sku="SKU-000789",
                name="Servidor Cloud Pro",
                category="Infraestructura",
                unit_price=Decimal("250.00"),
                unit_cost=Decimal("150.00"),
                is_active=True
            )
            session.add(prod)

        sale_num = int(uuid4().int % 900000 + 100000)
        sale = SaleModel(
            id=uuid4(),
            company_id=COMPANY_A_ID,
            branch_id=BRANCH_A_ID,
            customer_id=CUSTOMER_A_ID,
            series="F001",
            number=sale_num,
            status="confirmed",
            currency="USD",
            subtotal_amount=Decimal("250.00"),
            tax_amount=Decimal("45.00"),
            total_amount=Decimal("295.00"),
            paid_amount=Decimal("200.00")
        )
        session.add(sale)

        line = SaleLineModel(
            id=uuid4(),
            company_id=COMPANY_A_ID,
            sale_id=sale.id,
            product_id=PRODUCT_A_ID,
            line_number=1,
            sku_snapshot="SKU-000789",
            name_snapshot="Servidor Cloud Pro",
            unit_of_measure_snapshot="NIU",
            quantity=Decimal("1.00"),
            unit_price_snapshot=Decimal("250.00"),
            unit_cost_snapshot=Decimal("150.00"),
            discount_amount=Decimal("0.00"),
            tax_rate_snapshot=Decimal("0.18"),
            subtotal_amount=Decimal("250.00"),
            tax_amount=Decimal("45.00"),
            total_amount=Decimal("295.00")
        )
        session.add(line)

        wh = session.execute(
            select(WarehouseModel).where(WarehouseModel.id == WAREHOUSE_A_ID)
        ).scalar_one_or_none()
        if not wh:
            wh = WarehouseModel(
                id=WAREHOUSE_A_ID,
                company_id=COMPANY_A_ID,
                branch_id=BRANCH_A_ID,
                code="ALM-01",
                name="Almacen Central Lima",
                is_active=True
            )
            session.add(wh)

        kardex = KardexMovementModel(
            id=uuid4(),
            company_id=COMPANY_A_ID,
            warehouse_id=WAREHOUSE_A_ID,
            product_id=PRODUCT_A_ID,
            movement_type="INFLOW_PURCHASE",
            reference_doc_type="PURCHASE_ORDER",
            reference_doc_id=uuid4(),
            quantity_delta=Decimal("50.00"),
            resulting_physical_qty=Decimal("100.00"),
            unit_cost=Decimal("150.00"),
            moving_average_cost=Decimal("150.00"),
            occurred_at=datetime.now(timezone.utc)
        )
        session.add(kardex)

        susp_user = session.execute(
            select(UserModel).where(UserModel.id == SUSPENDED_USER_ID)
        ).scalar_one_or_none()
        if not susp_user:
            susp_user = UserModel(
                id=SUSPENDED_USER_ID,
                company_id=COMPANY_A_ID,
                email="suspended@salesia.enterprise",
                hashed_password="mocked_hash",
                role="sales",
                is_active=False
            )
            session.add(susp_user)

        session.commit()
    yield

@pytest.fixture
def auth_headers_a():
    identity = create_durable_identity(
        company_id=COMPANY_A_ID,
        user_id=USER_A_ID,
        role="company_admin",
        email="admin@salesia.enterprise"
    )
    return identity["headers"]

@pytest.fixture
def auth_headers_b():
    identity = create_durable_identity(
        company_id=COMPANY_B_ID,
        user_id=uuid4(),
        role="company_admin",
        email="admin@corp-b.enterprise"
    )
    return identity["headers"]

@pytest.fixture
def auth_headers_suspended():
    identity = create_durable_identity(
        company_id=COMPANY_A_ID,
        user_id=SUSPENDED_USER_ID,
        role="sales",
        email="suspended@salesia.enterprise"
    )
    return identity["headers"]

@pytest.mark.asyncio
async def test_rpt_001_catalog_unambiguous_and_filtering(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.get("/api/v1/reports", headers=auth_headers_a)
        assert res.status_code == 200
        data = res.json()
        assert data["count"] == 6
        for item in data["items"]:
            assert "code" in item
            assert "name" in item
            assert "area" in item
            assert "objective" in item
            assert "description" in item
            assert "formats" in item
            assert "xlsx" in item["formats"]
            assert "allowed_roles" in item

        res_comm = await client.get("/api/v1/reports?area=commercial", headers=auth_headers_a)
        assert res_comm.status_code == 200
        for item in res_comm.json()["items"]:
            assert item["area"] == "commercial"

        res_xlsx = await client.get("/api/v1/reports?format=xlsx", headers=auth_headers_a)
        assert res_xlsx.status_code == 200
        assert res_xlsx.json()["count"] == 6

@pytest.mark.asyncio
async def test_rpt_002_sales_detail_lines_and_no_duplication(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.post(
            "/api/v1/report-runs",
            json={"report_type": "REP_SALES_DETAIL", "format": "json"},
            headers=auth_headers_a
        )
        assert res.status_code == 202
        run_id = res.json()["run_id"]
        dl = await client.get(f"/api/v1/report-runs/{run_id}/download", headers=auth_headers_a)
        assert dl.status_code == 200
        payload = dl.json()
        assert "columns" in payload
        assert "line_number" in payload["columns"]
        assert "product_id" in payload["columns"]
        assert "total_amount" in payload["columns"]
        assert len(payload["rows"]) >= 1

@pytest.mark.asyncio
async def test_rpt_003_statistical_distribution_reproducible(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.post(
            "/api/v1/report-runs",
            json={"report_type": "REP_STATISTICAL_DISTRIBUTION", "format": "json"},
            headers=auth_headers_a
        )
        assert res.status_code == 202
        run_id = res.json()["run_id"]
        dl = await client.get(f"/api/v1/report-runs/{run_id}/download", headers=auth_headers_a)
        assert dl.status_code == 200
        payload = dl.json()
        metrics = {r["metric"]: r["value"] for r in payload["rows"]}
        assert "sample_size_n" in metrics
        assert "mean_ticket" in metrics
        assert "median_ticket" in metrics
        assert "variance" in metrics
        assert "standard_deviation" in metrics
        methods = {r["metric"]: r.get("method") for r in payload["rows"]}
        assert methods["mean_ticket"] == "sample_mean"

@pytest.mark.asyncio
async def test_rpt_004_crm_catalog_monetary_homogeneity(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.post(
            "/api/v1/report-runs",
            json={"report_type": "REP_CRM_CATALOG", "format": "json", "parameters": {"currency": "USD"}},
            headers=auth_headers_a
        )
        assert res.status_code == 202
        run_id = res.json()["run_id"]
        dl = await client.get(f"/api/v1/report-runs/{run_id}/download", headers=auth_headers_a)
        assert dl.status_code == 200
        payload = dl.json()
        for r in payload["rows"]:
            assert r["currency"] == "USD"
            assert "rank" in r

@pytest.mark.asyncio
async def test_rpt_005_sales_reps_commissions_and_refunds(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.post(
            "/api/v1/report-runs",
            json={"report_type": "REP_SALES_REPS_COMMISSIONS", "format": "json"},
            headers=auth_headers_a
        )
        assert res.status_code == 202
        run_id = res.json()["run_id"]
        dl = await client.get(f"/api/v1/report-runs/{run_id}/download", headers=auth_headers_a)
        assert dl.status_code == 200
        payload = dl.json()
        assert len(payload["rows"]) >= 1
        first = payload["rows"][0]
        assert "gross_amount" in first
        assert "returns_deductions" in first
        assert "net_amount" in first
        assert "payable_commission" in first

@pytest.mark.asyncio
async def test_rpt_006_kardex_inventory_balance_closure(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.post(
            "/api/v1/report-runs",
            json={"report_type": "REP_KARDEX_VALUATION", "format": "json"},
            headers=auth_headers_a
        )
        assert res.status_code == 202
        run_id = res.json()["run_id"]
        dl = await client.get(f"/api/v1/report-runs/{run_id}/download", headers=auth_headers_a)
        assert dl.status_code == 200
        payload = dl.json()
        assert len(payload["rows"]) >= 1
        first = payload["rows"][0]
        opening = Decimal(first["opening_balance"])
        inflow = Decimal(first["inflow_qty"])
        outflow = Decimal(first["outflow_qty"])
        closing = Decimal(first["closing_balance"])
        assert closing == opening + inflow - outflow

@pytest.mark.asyncio
async def test_rpt_007_cash_closure_receivables_no_fake_sales(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.post(
            "/api/v1/report-runs",
            json={"report_type": "REP_CASH_CLOSURE_RECEIVABLES", "format": "json"},
            headers=auth_headers_a
        )
        assert res.status_code == 202
        run_id = res.json()["run_id"]
        dl = await client.get(f"/api/v1/report-runs/{run_id}/download", headers=auth_headers_a)
        assert dl.status_code == 200
        payload = dl.json()
        assert len(payload["rows"]) >= 1
        first = payload["rows"][0]
        assert "sale_total" in first
        assert "payments_collected" in first
        assert "pending_receivable" in first
        assert "reconciliation_status" in first

@pytest.mark.asyncio
async def test_rpt_008_preview_sampling_and_limits(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        prev = await client.get(
            "/api/v1/reports/preview?report_type=REP_SALES_DETAIL&limit=10",
            headers=auth_headers_a
        )
        assert prev.status_code == 200
        data = prev.json()
        assert data["report_type"] == "REP_SALES_DETAIL"
        assert "columns" in data
        assert "estimated_total_rows" in data
        assert "estimated_file_size_kb" in data
        assert isinstance(data["preview_sample"], list)

@pytest.mark.asyncio
async def test_rpt_009_async_generation_and_cancellation(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        run_res = await client.post(
            "/api/v1/report-runs",
            json={"report_type": "REP_SALES_DETAIL", "format": "csv", "async_execution": True},
            headers=auth_headers_a
        )
        assert run_res.status_code == 202
        run_id = run_res.json()["run_id"]
        status_res = await client.get(f"/api/v1/report-runs/{run_id}", headers=auth_headers_a)
        assert status_res.status_code == 200
        assert status_res.json()["status"] in ("queued", "processing", "completed")

        with sync_session_factory() as session:
            art = session.execute(
                select(AnalyticsArtifact).where(AnalyticsArtifact.id == UUID(run_id))
            ).scalar_one()
            d = json.loads(art.data_json)
            d["status"] = "processing"
            art.status = "processing"
            art.data_json = json.dumps(d)
            session.commit()

        cancel_res = await client.post(f"/api/v1/report-runs/{run_id}/cancel", headers=auth_headers_a)
        assert cancel_res.status_code == 200
        assert cancel_res.json()["status"] == "cancelled"

@pytest.mark.asyncio
async def test_rpt_010_snapshot_immutability_and_hash(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        run_res = await client.post(
            "/api/v1/report-runs",
            json={"report_type": "REP_STATISTICAL_DISTRIBUTION", "format": "csv"},
            headers=auth_headers_a
        )
        assert run_res.status_code == 202
        run_data = run_res.json()
        assert run_data["data_hash"].startswith("sha256-")
        run_id = run_data["run_id"]

        get_res = await client.get(f"/api/v1/report-runs/{run_id}", headers=auth_headers_a)
        assert get_res.status_code == 200
        assert get_res.json()["data_hash"] == run_data["data_hash"]

@pytest.mark.asyncio
async def test_rpt_011_pdf_enterprise_layout(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.post(
            "/api/v1/report-runs",
            json={"report_type": "REP_SALES_DETAIL", "format": "pdf"},
            headers=auth_headers_a
        )
        assert res.status_code == 202
        run_id = res.json()["run_id"]
        dl = await client.get(f"/api/v1/report-runs/{run_id}/download", headers=auth_headers_a)
        assert dl.status_code == 200
        assert dl.headers["content-type"] == "application/pdf"
        assert dl.content.startswith(b"%PDF")
        assert b"%%EOF" in dl.content
        assert b"/Type /Page" in dl.content

@pytest.mark.asyncio
async def test_rpt_012_csv_formula_injection_defense(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        with sync_session_factory() as session:
            existing = session.scalar(select(CustomerModel).where(CustomerModel.company_id == COMPANY_A_ID, CustomerModel.code == "CLI-INJ"))
            if not existing:
                unique_doc = f"20{uuid4().int % 900000000 + 100000000}"
                malicious_cust = CustomerModel(
                    id=uuid4(),
                    company_id=COMPANY_A_ID,
                    code="CLI-INJ",
                    doc_type="RUC",
                    doc_number=unique_doc,
                    legal_name="=cmd|' /C calc'!A0",
                    credit_limit=Decimal("100.00"),
                    current_balance=Decimal("0.00"),
                    is_active=True
                )
                session.add(malicious_cust)
                session.commit()

        res = await client.post(
            "/api/v1/report-runs",
            json={"report_type": "REP_CRM_CATALOG", "format": "csv"},
            headers=auth_headers_a
        )
        assert res.status_code == 202
        run_id = res.json()["run_id"]
        dl = await client.get(f"/api/v1/report-runs/{run_id}/download", headers=auth_headers_a)
        assert dl.status_code == 200
        assert dl.text.startswith("\ufeff")
        assert "'=cmd|' /C calc'!A0" in dl.text
        assert "\n=cmd|" not in dl.text

@pytest.mark.asyncio
async def test_rpt_013_xlsx_typed_and_leading_zeros(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.post(
            "/api/v1/report-runs",
            json={"report_type": "REP_CRM_CATALOG", "format": "xlsx"},
            headers=auth_headers_a
        )
        assert res.status_code == 202
        run_id = res.json()["run_id"]
        dl = await client.get(f"/api/v1/report-runs/{run_id}/download", headers=auth_headers_a)
        assert dl.status_code == 200
        assert "spreadsheetml" in dl.headers["content-type"]
        wb = load_workbook(io.BytesIO(dl.content))
        sheet = wb.active
        assert sheet.title == "Reporte"
        found_leading_zero = False
        for row in sheet.iter_rows(values_only=True):
            for cell_val in row:
                if isinstance(cell_val, str) and cell_val.startswith("000"):
                    found_leading_zero = True
        assert found_leading_zero is True

@pytest.mark.asyncio
async def test_rpt_014_print_view_and_annulled_status(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        branding_res = await client.get("/api/v1/reports/branding", headers=auth_headers_a)
        assert branding_res.status_code == 200
        branding = branding_res.json()
        assert "print_layout" in branding
        assert branding["print_layout"]["page_size"] == "A4"

        res = await client.post(
            "/api/v1/report-runs",
            json={"report_type": "REP_SALES_DETAIL", "format": "pdf", "parameters": {"status": "cancelled"}},
            headers=auth_headers_a
        )
        assert res.status_code == 202
        run_id = res.json()["run_id"]
        dl = await client.get(f"/api/v1/report-runs/{run_id}/download", headers=auth_headers_a)
        assert dl.status_code == 200
        assert dl.content.startswith(b"%PDF")

@pytest.mark.asyncio
async def test_rpt_015_authorized_download_and_revocation(auth_headers_a, auth_headers_suspended):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.post(
            "/api/v1/report-runs",
            json={"report_type": "REP_SALES_DETAIL", "format": "csv"},
            headers=auth_headers_a
        )
        assert res.status_code == 202
        run_id = res.json()["run_id"]

        susp_dl = await client.get(f"/api/v1/report-runs/{run_id}/download", headers=auth_headers_suspended)
        assert susp_dl.status_code == 403

        with sync_session_factory() as session:
            art = session.execute(
                select(AnalyticsArtifact).where(AnalyticsArtifact.id == UUID(run_id))
            ).scalar_one()
            d = json.loads(art.data_json)
            d["is_revoked"] = True
            art.status = "revoked"
            art.data_json = json.dumps(d)
            session.commit()

        revoked_dl = await client.get(f"/api/v1/report-runs/{run_id}/download", headers=auth_headers_a)
        assert revoked_dl.status_code == 403

@pytest.mark.asyncio
async def test_rpt_016_safety_limits_enforced(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        overlimit = await client.post(
            "/api/v1/report-runs",
            json={"report_type": "REP_SALES_DETAIL", "limit": 999999},
            headers=auth_headers_a
        )
        assert overlimit.status_code in (400, 422)

@pytest.mark.asyncio
async def test_rpt_017_schedule_persistence_and_execution(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        create_res = await client.post(
            "/api/v1/report-schedules",
            json={
                "report_type": "REP_SALES_DETAIL",
                "cron_expression": "0 8 * * 1-5",
                "format": "xlsx",
                "recipients": ["auditoria@salesia.enterprise"],
                "timezone": "America/Lima"
            },
            headers=auth_headers_a
        )
        assert create_res.status_code == 201
        sched_id = create_res.json()["schedule_id"]

        exec_res = await client.post(f"/api/v1/report-schedules/{sched_id}/execute", headers=auth_headers_a)
        assert exec_res.status_code == 200
        assert exec_res.json()["status"] == "completed"

        del_res = await client.delete(f"/api/v1/report-schedules/{sched_id}", headers=auth_headers_a)
        assert del_res.status_code == 200

@pytest.mark.asyncio
async def test_rpt_018_channel_delivery_authorization(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        bad_recipient = await client.post(
            "/api/v1/report-schedules",
            json={
                "report_type": "REP_SALES_DETAIL",
                "cron_expression": "0 8 * * *",
                "format": "csv",
                "recipients": ["attacker@external-spam-domain.com"]
            },
            headers=auth_headers_a
        )
        assert bad_recipient.status_code in (400, 422)

@pytest.mark.asyncio
async def test_rpt_019_history_and_expiration_http_410(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.post(
            "/api/v1/report-runs",
            json={"report_type": "REP_SALES_DETAIL", "format": "csv"},
            headers=auth_headers_a
        )
        run_id = res.json()["run_id"]

        with sync_session_factory() as session:
            art = session.execute(
                select(AnalyticsArtifact).where(AnalyticsArtifact.id == UUID(run_id))
            ).scalar_one()
            d = json.loads(art.data_json)
            expired_time = (datetime.now(timezone.utc) - timedelta(minutes=20)).isoformat()
            d["expires_at"] = expired_time
            art.data_json = json.dumps(d)
            session.commit()

        expired_dl = await client.get(f"/api/v1/report-runs/{run_id}/download", headers=auth_headers_a)
        assert expired_dl.status_code == 410

@pytest.mark.asyncio
async def test_rpt_020_renderer_failure_cleanup_and_retry(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.post(
            "/api/v1/report-runs",
            json={"report_type": "REP_SALES_DETAIL", "format": "csv"},
            headers=auth_headers_a
        )
        run_id = res.json()["run_id"]

        with sync_session_factory() as session:
            art = session.execute(
                select(AnalyticsArtifact).where(AnalyticsArtifact.id == UUID(run_id))
            ).scalar_one()
            d = json.loads(art.data_json)
            d["status"] = "failed"
            d["error_cause"] = "simulated_renderer_failure"
            art.status = "failed"
            art.data_json = json.dumps(d)
            session.commit()

        retry_res = await client.post(f"/api/v1/report-runs/{run_id}/retry", headers=auth_headers_a)
        assert retry_res.status_code == 200
        assert retry_res.json()["status"] == "completed"

@pytest.mark.asyncio
async def test_rpt_021_report_comparison(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        r1 = await client.post("/api/v1/report-runs", json={"report_type": "REP_SALES_DETAIL"}, headers=auth_headers_a)
        r2 = await client.post("/api/v1/report-runs", json={"report_type": "REP_SALES_DETAIL"}, headers=auth_headers_a)
        id1 = r1.json()["run_id"]
        id2 = r2.json()["run_id"]
        cmp_res = await client.post(
            "/api/v1/reports/compare",
            json={"run_id_1": id1, "run_id_2": id2},
            headers=auth_headers_a
        )
        assert cmp_res.status_code == 200
        data = cmp_res.json()
        assert "row_count_delta" in data
        assert "is_identical" in data
        assert "cutoff_delta_seconds" in data

@pytest.mark.asyncio
async def test_rpt_022_company_branding_safe_sanitization(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.get("/api/v1/reports/branding", headers=auth_headers_a)
        assert res.status_code == 200
        data = res.json()
        assert "<script>" not in data["legal_name"]
        assert "safe_logo_url" in data
        assert "template_version" in data

@pytest.mark.asyncio
async def test_api_082_to_085_contracts_and_cross_tenant_isolation(auth_headers_a, auth_headers_b):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        cat = await client.get("/api/v1/reports", headers=auth_headers_a)
        assert cat.status_code == 200

        run = await client.post(
            "/api/v1/report-runs",
            json={"report_type": "REP_SALES_DETAIL", "format": "xlsx"},
            headers=auth_headers_a
        )
        assert run.status_code == 202
        run_id = run.json()["run_id"]

        status = await client.get(f"/api/v1/report-runs/{run_id}", headers=auth_headers_a)
        assert status.status_code == 200
        assert status.json()["status"] == "completed"

        tenant_b_status = await client.get(f"/api/v1/report-runs/{run_id}", headers=auth_headers_b)
        assert tenant_b_status.status_code == 404

        tenant_b_dl = await client.get(f"/api/v1/report-runs/{run_id}/download", headers=auth_headers_b)
        assert tenant_b_dl.status_code in (404, 403)

        dl = await client.get(f"/api/v1/report-runs/{run_id}/download", headers=auth_headers_a)
        assert dl.status_code == 200

@pytest.mark.asyncio
async def test_reg_18_all_formats_open_cleanly(auth_headers_a):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        for fmt in ["csv", "xlsx", "pdf", "json"]:
            run = await client.post(
                "/api/v1/report-runs",
                json={"report_type": "REP_SALES_DETAIL", "format": fmt},
                headers=auth_headers_a
            )
            assert run.status_code == 202
            run_id = run.json()["run_id"]
            dl = await client.get(f"/api/v1/report-runs/{run_id}/download", headers=auth_headers_a)
            assert dl.status_code == 200
            if fmt == "csv":
                assert dl.text.startswith("\ufeff")
            elif fmt == "xlsx":
                wb = load_workbook(io.BytesIO(dl.content))
                assert wb.active is not None
            elif fmt == "pdf":
                assert dl.content.startswith(b"%PDF")
                assert b"%%EOF" in dl.content
            elif fmt == "json":
                assert isinstance(dl.json(), dict)
